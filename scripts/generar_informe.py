"""Regenera el informe usando los resultados guardados y las figuras del proyecto."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parents[1]
DATA=json.loads((BASE/'evidencias/automatizadas/RESUMEN_BUILD.json').read_text())
assert DATA['build_success'] and DATA['core_failures']==0 and DATA['core_errors']==0
DEST=BASE/'docs/Informe_FrioSeguro_OT.docx'
doc=Document(DEST)
for el in list(doc._element.body):
 if el.tag!=qn('w:sectPr'):doc._element.body.remove(el)
sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(7)
 if name=='Normal':st.font.size=Pt(10.5);st.paragraph_format.line_spacing=1.1
for name,size in [('Title',29),('Subtitle',13),('Heading 1',19),('Heading 2',12.5)]:doc.styles[name].font.size=Pt(size)
for st in doc.styles:
 if st.type==1:
  pr=st.element.find(qn('w:pPr'))
  if pr is not None:
   for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
for part in [sec.header,sec.footer]:
 for item in part.paragraphs:item.clear()
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('FríoSeguro OT  |  ').font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def p(text='',bold=False):
 a=doc.add_paragraph();a.add_run(text).bold=bold;return a
def h(text):return doc.add_heading(text,1)
def sub(text):return doc.add_heading(text,2)
def page():doc.add_page_break()
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,value in enumerate(headers):t.rows[0].cells[i].text=value
 for row in rows:
  cells=t.add_row().cells
  for i,value in enumerate(row):cells[i].text=str(value)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if ri==0:trpr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'5');b.set(qn('w:color'),'D9D9D9');borders.append(b)
   pr.append(borders);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'173B50' if ri==0 else ('F0F5F7' if ri%2==0 else 'FFFFFF'));pr.append(sh)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margins.append(e)
   pr.append(margins)
   for paragraph in c.paragraphs:
    paragraph.paragraph_format.space_after=Pt(2);paragraph.paragraph_format.line_spacing=1.03
    for run in paragraph.runs:run.font.size=Pt(9.4);run.bold=ri==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
 p();return t
def picture(filename,width,alt):
 a=doc.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER
 im=a.add_run().add_picture(str(BASE/'docs/figuras'/filename),width=Inches(width));im._inline.docPr.set('descr',alt)
def caption(text):
 a=p(text)
 for r in a.runs:r.italic=True;r.font.size=Pt(9)

p('INACAP  ·  TI3042',True)
doc.add_paragraph('FríoSeguro OT',style='Title')
doc.add_paragraph('Aplicación Android para monitoreo y control de una cámara de refrigeración simulada',style='Subtitle')
p('Néstor Urdaneta',True);p('Aplicaciones Móviles para IoT · Unidad 2')
p('Evaluación sumativa individual · Ponderación 25 % · Plazo de trabajo 6 semanas')
p('Versión 1.1 · Paquete final del 25 de octubre de 2026')
sub('Propósito del proyecto')
p('Presento FríoSeguro OT, un prototipo que conecta dos dispositivos Android para observar temperatura, puerta y compresor de una cámara simulada y solicitar cambios de control con confirmación. La versión sumativa incorpora autenticación local en SQLite, conservación de lecturas, consulta con filtros y exportación CSV. La comunicación utiliza Wi-Fi local con TLS y verificación de la identidad del Equipo.')
p('Una misma APK ofrece los roles Equipo y Operador. El Equipo simula sensores y actuador; el Operador monitorea y envía órdenes. Cada instalación conserva sus propios datos, incluso después de desconectar o cerrar sesión. La solución funciona sin un servicio de Internet.')
sub('Alcance de la entrega')
p('Se incluyen código Kotlin, APK de desarrollo, diseños de pantallas, esquema SQLite, medidas de seguridad, resultados de verificación y un protocolo de validación en dos teléfonos. Los umbrales y las consignas son didácticos. La app no opera un refrigerador, un PLC ni un proceso industrial real.')
p('La demostración inalámbrica en dos teléfonos y la aprobación del docente deben obtenerse durante la evaluación. Las comprobaciones del emulador y del computador se identifican por separado y no acreditan alcance, consumo ni estabilidad de una Wi-Fi física.')
sub('Contenido del informe')
p('Las páginas siguientes explican la selección de herramientas, las pantallas, el almacenamiento, el enlace seguro, las pruebas y la correspondencia con la pauta. El repositorio contiene las instrucciones para reproducir las comprobaciones y entregar el enlace en el AAI.')

page();h('Herramientas y decisiones del proyecto')
p('El problema consiste en detectar a distancia condiciones anómalas y comprobar si una orden se aplicó. El control local mantiene un comportamiento definido cuando se interrumpe la sesión. Se conserva la temática FríoSeguro OT del avance anterior; su continuidad con la Unidad 1 debe ser confirmada por el docente.')
sub('Desarrollo Android')
p('La app está escrita en Kotlin para Android Studio. AppCompat y Material Components aportan campos, botones, slider, radio buttons, switch y navegación. ViewModel y LiveData separan la presentación del estado. Los ejecutores atienden red, autenticación y almacenamiento fuera del hilo de interfaz.')
p('La configuración fija Kotlin 2.1.10, Android Gradle Plugin 8.9.2, Gradle 8.11.1, JDK 17, compile y target SDK 35 y mínimo Android 8.0 API 26. SQLiteOpenHelper implementa el esquema acotado de esta actividad; Room sería una evolución apropiada para un modelo más amplio [1].')
sub('Elección del almacenamiento')
table(['Opción','Seguridad y escala','Facilidad y decisión'],[
 ['SQLite local','Sandbox Android y payloads cifrados. Una cuenta y hasta 10.000 lecturas por instalación.','Seleccionada. Funciona sin Internet y permite demostrar autenticación, guardado y consulta.'],
 ['Firebase','Exige reglas, proyecto y autenticación del servicio. Facilita la gestión centralizada.','Requiere configurar servicios externos. Alternativa para varias ubicaciones.'],
 ['API con MySQL','Exige backend HTTPS y autorización en servidor. Permite datos centralizados.','Mayor operación. Android consultaría una API, no MySQL directamente.']],[1.05,2.65,2.95])
p('La opción implementada es SQLite. JSON versionado y CSV con unidades y tiempos documentados permiten intercambiar datos con hojas de cálculo o herramientas externas sin incorporar un servidor a la demostración.')
sub('Elección de la comunicación')
table(['Criterio','Justificación de Wi-Fi'],[
 ['Distancia','Adecuada para una sala y para ensayar una pared. El alcance real se medirá con los teléfonos.'],
 ['Energía','Las apps en primer plano y una muestra por segundo priorizan la observación. No se promete la autonomía de un sensor BLE dormido.'],
 ['Estabilidad','TCP ordena y retransmite. La app detecta datos vencidos y confirma órdenes; una pérdida exige reconectar.']],[1.05,5.6])
p('Wi-Fi con TLS sobre TCP evita depender de Internet o un broker MQTT. Bluetooth es una alternativa para corta distancia y bajo consumo. Se utiliza un router local o un hotspot de un tercer teléfono, sin aislamiento de clientes.')

page();h('Pantallas principales y navegación')
p('El acceso exige una cuenta local. Después del login, cinco destinos organizan Conexión, Monitoreo, Activador, Registro e Historial. Los controles incluyen etiquetas y estados con texto, áreas táctiles de 48 dp y contenedores desplazables.')
picture('Mockup_pantallas.png',6.2,'Bosquejos de Login, Conexión, Monitoreo y Activador con sus controles principales')
caption('Figura 1. Bosquejos de diseño. Los valores ilustran estados de la interfaz y no corresponden a una medición física.')
p('El Activador habilita órdenes para un Operador autenticado, con enlace reciente y permiso remoto del Equipo. La interfaz muestra la orden pendiente y su confirmación. Login y Conexión protegen la visualización de secretos frente a capturas de pantalla.')

page();h('Datos persistentes y consulta local')
p('SQLite conserva la cuenta y las lecturas de cada instalación. Historial permite consultar aun sin Wi-Fi. Registro muestra eventos y métricas de la ejecución; ambos tienen finalidades distintas.')
picture('Mockup_historial.png',6.25,'Bosquejos de Historial con filtros y CSV y de Registro con métricas de ejecución')
caption('Figura 2. Historial incorpora origen, periodo y alarma; Registro informa TLS, ACK y errores de almacenamiento.')
sub('Cuenta y protección de datos')
p('account almacena un único sobre cifrado con usuario, sal aleatoria y hash PBKDF2 HMAC SHA256 de 210.000 iteraciones. La contraseña no se guarda. Android Keystore protege la clave AES 256 GCM. SQLite conserva fallos y el bloqueo de 30 segundos tras cinco intentos incorrectos [3].')
p('AuthStore migra las preferencias de la versión 1.0 cuando la instalación conserva su identidad de firma. Verifica la escritura en SQLite antes de retirar la copia anterior. Desinstalar elimina datos y claves; la APK final de demostración puede exigir reinstalar por su nueva firma.')
sub('Captura y consulta')
p('telemetry guarda ID, hora de captura y recepción, origen LOCAL o REMOTO, alarma y payload JSON versión 2 cifrado. Temperatura, consigna, modo, compresor, puerta y número de muestra están dentro del payload. Los metadatos de índice quedan visibles y se autentican como datos asociados; la base completa no usa SQLCipher.')
p('Cada guardado transaccional conserva hasta 10.000 lecturas. Los índices por tiempo, origen y alarma y los argumentos enlazados permiten filtrar sin concatenar valores al SQL. La pantalla muestra las 50 coincidencias recientes y actualiza cada cinco segundos. El CSV exporta todas las coincidencias conservadas en orden cronológico, UTF-8, fechas UTC ISO 8601, grados Celsius, punto decimal y booleanos 0/1.')

page();h('Arquitectura del enlace y del control')
picture('Arquitectura.png',6.45,'Dos Android enlazados por Wi-Fi TLS con almacenes SQLite privados independientes')
caption('Figura 3. El enlace comparte telemetría y órdenes; las cuentas permanecen en cada Android.')
sub('Emparejamiento y canal seguro')
p('El Equipo genera una identidad RSA no exportable mediante Android Keystore y abre el puerto 8842. Su código FS2 reúne un secreto aleatorio de 128 bits y la huella SHA256 del certificado. El Operador recibe el código completo por un medio de confianza y comprueba huella, vigencia y firma antes de completar TLS [2, 3].')
p('Se admiten TLS 1.2 o 1.3 con suites AEAD e intercambio efímero. No hay ruta en texto plano. La fijación de identidad corresponde al par local emparejado manualmente; si cambia el certificado debe compartirse el nuevo código. No es una política de pinning de un servicio remoto permanente.')
p('La clave autoriza DIGEST_NONE para firmar el resumen que calcula la pila TLS [8]. Después de TLS, SecureChannel verifica la PSK con desafíos HMAC y deriva claves por dirección con HKDF SHA256 [7]. Su sobre AES GCM aplica secuencias y límites de trama. El protocolo académico adicional evita aceptar un cliente que solo conozca el certificado público.')
sub('Mensajes y recuperación')
p('El Equipo emite telemetría cada segundo con schemaVersion 2 y hora de captura. El Operador mide RTT con ping y distingue datos recientes de datos vencidos a los tres segundos. La lectura del socket expira a los siete segundos. Los intentos iniciales son acotados y una sesión perdida requiere reconexión manual.')
p('Las órdenes llevan UUID y tiempo. El Equipo valida permiso remoto, vigencia de 60 segundos, duplicados, frecuencia, modo y valores finitos. Se admite una orden pendiente y se espera ACK por cuatro segundos. No se reenvía automáticamente una orden sin confirmar porque podría haberse aplicado.')
p('AUTO utiliza histéresis de 0,5 °C. Ante pérdida del enlace, el Equipo vuelve a AUTO con consigna 4 °C mientras la simulación siga activa. Al abandonar la app se cierra la conexión; tras cinco minutos sin interacción se cierra la sesión. El historial permanece disponible para el siguiente ingreso.')

page();h('Seguridad IoT y lineamientos OT')
p('Las medidas se relacionan con los objetivos de seguridad y privacidad de ISO IEC 27400, el enfoque de gestión de riesgos de ISO IEC 27001 y la prioridad OT de operación segura y disponibilidad de NIST SP 800 82 revisión 3 [4, 5, 6]. Esta relación no constituye certificación ISO ni habilita uso industrial.')
table(['Riesgo','Medida implementada','Verificación o límite'],[
 ['Acceso local indebido','SQLite, PBKDF2, AES GCM con Keystore y bloqueo persistente.','Login válido e inválido, reapertura y cuenta cifrada. Una cuenta por instalación.'],
 ['Suplantación o escucha','TLS con huella transferida fuera de banda, PSK y claves por dirección.','Rechazo de huella y código incorrectos. La confianza depende de compartir bien el FS2.'],
 ['Alteración o repetición','AEAD, secuencias, IDs de orden, vigencia y límites de trama.','Pruebas de tag, replay, formato y órdenes inválidas.'],
 ['Control no autorizado','Login, permiso remoto del Equipo y enlace reciente.','Desactivar permiso y comprobar el bloqueo del Activador.'],
 ['Pérdida del enlace','Datos vencidos, timeout, ACK y retorno local a AUTO 4 °C.','Corte físico pendiente. La simulación exige proceso activo.'],
 ['Divulgación de datos','Sandbox privado, payload cifrado, sin contraseñas en logs y exportación explícita.','El CSV contiene lecturas legibles elegidas por el usuario. Metadatos de índice visibles.'],
 ['Saturación o fallo de disco','10.000 registros, cola de 64 y contadores de error.','Prueba de retención y errores visibles en Registro. No se promete captura sin pérdida.']],[1.3,2.6,2.75])
sub('Aplicación de los objetivos de seguridad')
p('Identidad y acceso se atienden con cuenta local, verificación del par y permiso por sesión. La confidencialidad e integridad se atienden en transporte y almacenamiento. La disponibilidad se aborda con control local y fallos detectables. La minimización evita recopilar ubicación, contactos o identificadores personales innecesarios.')
p('Para la demostración se propone una red de laboratorio separada, con contraseña del router, sin exposición del puerto a Internet y con acceso físico controlado a los teléfonos. Esa segmentación debe verificarse; la APK no la impone al router.')
p('Las bitácoras ayudan a reconstruir la sesión, pero no son inmutables. Una app Android en primer plano no sustituye un controlador de seguridad. Una futura integración física exige análisis de riesgos, gestión de claves y validación de actuadores apropiados al proceso.')

page();h('Verificación de la aplicación')
p('La verificación distingue el núcleo JVM, la instrumentación Android y los ensayos físicos. La reconstrucción del paquete final del 25 de octubre conserva la lógica Kotlin probada el 24 de octubre. Las capturas antiguas de la versión 1.0 están identificadas como antecedentes.')
table(['Comprobación','Resultado y alcance'],[
 ['Compilación y firma','APK 1.1 y APK de pruebas construidas con JDK 17 y SDK 35. Firma de desarrollo Android Debug y SHA256 del archivo verificados.'],
 ['Núcleo JVM',f"{DATA['core_tests']} pruebas aprobadas el 25 de octubre, sin fallos ni errores: planta, canal seguro y emparejamiento TLS."],
 ['Android API 28','12 casos aprobados el 24 de octubre en dos lotes: 9 de cuenta/historial y 3 de TLS/integración repetidos tras corregir la clave TLS.'],
 ['Persistencia entre procesos','seed y verify aprobaron en instrumentaciones separadas con am force-stop entre ambas. Persistieron cuenta y lectura REMOTO.'],
 ['Análisis estático',f"Android Lint: {DATA['lint_errors']} errores y {DATA['lint_warnings']} advertencias. Desglose en VERIFICACION.md."],
 ['Dos Android físicos','Pendiente. No se han medido alcance, consumo, RTT ni estabilidad de una Wi-Fi física.']],[1.65,5.0])
sub('Alcance de las comprobaciones')
p('Las pruebas JVM comprueban conexiones TLS 1.2 y 1.3 reales en loopback, rechazo de huella incorrecta, certificados inválidos y códigos antiguos. Cubren además manipulación de tramas, repetición, límites y comportamiento de la planta. Estas conexiones locales no representan ensayos inalámbricos.')
p('En Android se verificaron cuenta cifrada, bloqueo y migración; guardado y filtros por origen, periodo y alarma; coherencia CSV; entradas inválidas, metadatos alterados y retención. La retención prepara filas por SQL y comprueba el límite al guardar una lectura real; no equivale a varias horas de captura física.')
p('La integración usó dos modelos en un solo emulador. Login y navegación se activaron mediante eventos programáticos de vistas Android. Comprobó telemetría, alarma, ACK, consulta del Historial y retorno a AUTO. El resultado del 24 de octubre se conserva como resumen fechado; los logs y capturas temporales de esa versión no se incluyen. No se repitió la instrumentación Android el 25 de octubre.')
sub('Ensayos que deben completarse en clase')
p('El protocolo T01 a T23 y el registro CSV mantienen PENDIENTE los casos sin ejecución física. Deben anotarse dispositivos, versión Android, topología, distancia, duración, muestras, RTT, saltos, ACK, consulta SQLite y filas CSV. Una casilla vacía no acredita el criterio.')

page();h('Correspondencia con la evaluación')
p('La pauta de la fotografía suma 100 puntos. Esta matriz permite localizar la implementación y la evidencia de cada criterio. El docente asigna el puntaje al revisar la aplicación y la demostración.')
table(['Criterio','Puntos','Implementación y evidencia'],[
 ['2.1.1.1 Login Android','10','Cuenta y validación SQLite, contraseña oculta y bloqueo. AccountSecurityTest.'],
 ['2.1.1.2 UI Material','10','Monitoreo, Activador, navegación, slider, radio buttons y switch. Bosquejos y vistas Android.'],
 ['2.1.2.3 Conexión segura','15','Wi-Fi TLS y PSK, límites y métricas. Pruebas de canal; Wi-Fi física pendiente.'],
 ['2.1.2.4 Solución efectiva','10','Alarmas, histéresis, ACK, fallback e historial persistente.'],
 ['2.1.3.5 Seguridad ISO','10','Medidas relacionadas con ISO 27400 y 27001. SEGURIDAD.md y pruebas negativas.'],
 ['2.1.3.6 Lineamientos OT','10','Riesgos, disponibilidad local, propuesta de segmentación y límites documentados.'],
 ['2.1.4.7 Interconexión','15','APK con roles Equipo y Operador. Demostración pendiente en dos teléfonos.'],
 ['2.1.4.8 Estabilidad','20','Métricas y escenarios T11 a T15; mediciones físicas pendientes.'],
 ['Total','100','Cobertura implementada y protocolo de demostración preparado.']],[1.65,.55,4.45])
sub('Requisitos de la guía sumativa')
p('La autenticación usa la opción SQLite autorizada por la guía. Los datos capturados se almacenan y consultan mediante esa opción. DATOS_SQLITE.md compara seguridad, escalabilidad y facilidad frente a Firebase y API con MySQL y documenta esquema, transacciones y formato interoperable.')
p('El transporte incorpora TLS estándar. DISENO.md justifica Wi-Fi según distancia, energía y estabilidad. El informe incorpora Historial y Registro y un guion que incluye almacenamiento y consulta, además de autenticación y control.')
sub('Revisión del docente')
p('Presentar problemática, continuidad con Unidad 1, método de conexión, bosquejos, autenticación y esquema. Registrar la retroalimentación real y aplicar los ajustes indicados. La aprobación del diseño, la asistencia y la demostración presencial forman parte de la evaluación y deben obtenerse del docente.')

page();h('Instalación y demostración en clase')
sub('Preparar dos teléfonos')
p('Instalar apk/FrioSeguroOT.apk en dos Android 8.0 o posteriores. La APK final se firmó de nuevo. Si Android rechaza actualizar una demostración anterior por firma distinta, exportar antes su bitácora y desinstalarla; se eliminan sus datos y claves. Instalar la APK final y crear una cuenta en cada teléfono. La migración automática del código requiere conservar la misma identidad de firma.')
p('Conectar ambos a la misma Wi-Fi sin aislamiento de clientes y mantener las apps visibles. En B elegir Equipo, habilitar permiso remoto e iniciar. En A elegir Operador, copiar la IPv4 privada y el código FS2 completo de B por un medio de confianza y conectar. Usar fecha y hora automáticas.')
sub('Guion de ocho a diez minutos')
for text in [
 '1. Explicar el problema y justificar SQLite y Wi-Fi. Crear o utilizar la cuenta local y demostrar el rechazo de una contraseña incorrecta.',
 '2. Mostrar roles, dirección del Equipo, estado TLS y telemetría. Verificar puerta, temperatura, compresor y tendencia.',
 '3. Provocar calor o apertura de puerta en B. Mostrar la alarma en A. Enviar una orden MANUAL y comprobar su ACK y efecto en B.',
 '4. Abrir Historial, filtrar REMOTO y alarmas, consultar y exportar CSV. Comparar conteo y filas. Desconectar y comprobar la consulta sin red.',
 '5. Cerrar y abrir la app e ingresar otra vez. Verificar persistencia. Ensayar pérdida de Wi-Fi, datos vencidos y fallback AUTO con consigna 4 °C.',
 '6. Explicar Keystore, TLS, huella, PSK y límites. Mostrar pruebas y completar el registro físico con resultados observados.'
]:p(text)
sub('Entrega en el AAI')
p('La guía solicita un Bloc de notas con el enlace al código en GitHub. Subir Entrega_TI3042_Nestor_Urdaneta.txt en la tarea correcta antes del cierre y pulsar Enviar tarea para confirmar. Verificar el estado de entrega. No se reciben entregas por correo.')
p('Repositorio público: https://github.com/nestorperezurdaneta23-create/FrioSeguroOT')
p('El TXT mantiene ese enlace. El repositorio incluye APK, informe, fuentes y evidencias. La demostración y el registro de los teléfonos complementan la entrega del archivo.')

page();h('Referencias y documentos del proyecto')
p('La fuente de requisitos es la Guía de Aprendizaje Sumativa Unidad 2 TI3042 proporcionada por el docente. La pauta de la fotografía se conserva como matriz de ocho criterios y 100 puntos. La guía indica evaluación individual de 25 %, seis semanas y entrega en AAI mediante TXT con enlace a GitHub.')
refs=[
 ('1','Android Developers','Save data using SQLite','https://developer.android.com/training/data-storage/sqlite'),
 ('2','Android Developers','Security with network protocols y TLS','https://developer.android.com/privacy-and-security/security-ssl'),
 ('3','Android Developers','Android Keystore system','https://developer.android.com/privacy-and-security/keystore'),
 ('4','ISO','ISO IEC 27400 Security and privacy guidelines for IoT','https://www.iso.org/standard/44373.html'),
 ('5','ISO','ISO IEC 27001 Information security management systems','https://www.iso.org/standard/27001'),
 ('6','NIST','SP 800 82 Rev 3 Guide to Operational Technology Security','https://csrc.nist.gov/pubs/sp/800/82/r3/final'),
 ('7','IETF','RFC 5869 HMAC based Extract and Expand Key Derivation Function','https://www.rfc-editor.org/rfc/rfc5869'),
 ('8','Android Developers','KeyGenParameterSpec Builder y resúmenes autorizados para TLS','https://developer.android.com/reference/android/security/keystore/KeyGenParameterSpec.Builder')]
for n,author,title,url in refs:
 a=p(f'[{n}] {author}. {title}.',True);a.paragraph_format.space_after=Pt(3)
 a=p(url);a.paragraph_format.space_after=Pt(8)
 for run in a.runs:run.font.size=Pt(9.5)
sub('Archivos para revisar')
p('README.md explica instalación y compilación. DATOS_SQLITE.md documenta tablas, cifrado, consultas y CSV. DISENO.md presenta arquitectura y pantallas. SEGURIDAD.md relaciona riesgos, medidas, pruebas y límites. MATRIZ_PAUTA.md vincula cada criterio con código y evidencia.')
p('PRUEBAS_DISPOSITIVOS.md y Registro_pruebas_wifi.csv permiten registrar los ensayos físicos. VERIFICACION.md distingue resultados actuales y antecedentes. El proyecto incluye :core para lógica y protocolo JVM y :app para UI, SQLite, Keystore y TLS. Los tests instrumentados preparan o borran la cuenta y lecturas del dispositivo de prueba; se ejecutan en un dispositivo reservado para ello.')

doc.core_properties.title='FríoSeguro OT Aplicación Android para la evaluación sumativa'
doc.core_properties.subject='TI3042 Unidad 2 SQLite TLS y control IoT simulado'
doc.core_properties.author='Néstor Urdaneta';doc.core_properties.last_modified_by='Néstor Urdaneta'
doc.save(DEST);print(DEST)
